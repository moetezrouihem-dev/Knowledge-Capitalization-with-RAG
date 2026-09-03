import os
import re
import json
from pathlib import Path

import docx
import openpyxl
import pdfplumber
from bs4 import BeautifulSoup

SRC_ROOT = Path("docs")
OUT_ROOT = Path("corpus_text")
OUT_ROOT.mkdir(exist_ok=True)


def safe_name(p: Path) -> str:
    return re.sub(r"[^\w\-]+", "_", p.stem)[:80]


def extract_docx(path: Path) -> str:
    d = docx.Document(str(path))
    parts = []
    for para in d.paragraphs:
        if para.text.strip():
            style = para.style.name if para.style else ""
            prefix = "## " if "Heading" in style else ""
            parts.append(prefix + para.text.strip())
    for table in d.tables:
        parts.append("\n[TABLE]")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_xlsx(path: Path) -> str:
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        parts.append(f"\n## Sheet: {sheet_name}")
        for row in ws.iter_rows(values_only=True):
            vals = [str(v) for v in row if v is not None and str(v).strip()]
            if vals:
                parts.append(" | ".join(vals))
    return "\n".join(parts)


def extract_pdf(path: Path) -> str:
    parts = []
    with pdfplumber.open(str(path)) as pdf:
        for i, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            if text.strip():
                parts.append(f"\n## Page {i+1}\n{text}")
            for table in page.extract_tables() or []:
                parts.append("[TABLE]")
                for row in table:
                    parts.append(" | ".join(c or "" for c in row))
    return "\n".join(parts)


def _parse_js_objects(js_array: str):
   
    no_comments = re.sub(r"(?m)^[ \t]*//.*$", "", js_array)
    objects = []
    depth = 0
    start = None
    for i, ch in enumerate(no_comments):
        if ch == "{":
            if depth == 0:
                start = i
            depth += 1
        elif ch == "}":
            depth -= 1
            if depth == 0 and start is not None:
                objects.append(no_comments[start:i + 1])
                start = None

    entries = []
    pair_re = re.compile(
        r"(\w+)\s*:\s*(?:'((?:[^'\\]|\\.)*)'|\"((?:[^\"\\]|\\.)*)\")"
    )
    for obj in objects:
        d = {}
        for key, sq_val, dq_val in pair_re.findall(obj):
            val = sq_val if sq_val != "" or "'" in obj else dq_val
            val = (sq_val or dq_val).replace("\\'", "'").replace('\\"', '"')
            d[key] = val
        if d:
            entries.append(d)
    return entries


def extract_html(path: Path) -> str:
    raw = path.read_text(encoding="utf-8", errors="ignore")
    match = re.search(r"let\s+D\s*=\s*(\[.*?\]);", raw, re.DOTALL)
    parts = []
    if match:
        js_array = match.group(1)
        entries = _parse_js_objects(js_array)
        if entries:
            parts.append(f"## QMS Register ISO 9001 — {len(entries)} entries\n")
            for e in entries:
                line = " | ".join(f"{k}: {v}" for k, v in e.items())
                parts.append(line)
        else:
            parts.append("## QMS Register (raw extraction)\n" + js_array[:5000])
    soup = BeautifulSoup(raw, "lxml")
    visible = soup.get_text(separator="\n")
    visible = "\n".join(l.strip() for l in visible.splitlines() if l.strip())
    parts.append("\n## Visible page text\n" + visible)
    return "\n".join(parts)


def main():
    manifest = []
    for path in sorted(SRC_ROOT.rglob("*")):
        if path.is_dir():
            continue
        ext = path.suffix.lower()
        try:
            if ext == ".docx":
                text = extract_docx(path)
            elif ext == ".xlsx":
                text = extract_xlsx(path)
            elif ext == ".pdf":
                text = extract_pdf(path)
            elif ext == ".html":
                text = extract_html(path)
            else:
                continue
        except Exception as e:
            print(f"ERROR on {path}: {e}")
            continue

        category = path.parent.name if path.parent != SRC_ROOT else "Root"
        out_name = f"{safe_name(path)}.txt"
        out_path = OUT_ROOT / out_name
        header = f"SOURCE: {path.name}\nCATEGORY: {category}\n\n"
        out_path.write_text(header + text, encoding="utf-8")
        manifest.append({
            "source_file": path.name,
            "category": category,
            "extracted_file": out_name,
            "n_chars": len(text),
        })
        print(f"OK  {path.name}  ->  {out_name}  ({len(text)} chars)")

    Path("manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"\nTotal: {len(manifest)} documents extracted.")


if __name__ == "__main__":
    main()
